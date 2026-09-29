import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

def create_observability_tracking_excel(filename="docs/ISC_Form_Tracking_Observability_Services.xlsx"):
    wb = openpyxl.Workbook()

    # Create sheets
    ws_summary = wb.active
    ws_summary.title = "01_TONG_QUAN"
    ws_tracking = wb.create_sheet(title="02_TRACKING_SERVICES")
    ws_catalog = wb.create_sheet(title="03_DANH_MUC_CHUAN")

    # =========================================================================
    # STYLES & PALETTE
    # =========================================================================
    font_family = "Segoe UI"

    title_font = Font(name=font_family, size=16, bold=True, color="FFFFFF")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="E2E8F0")
    section_font = Font(name=font_family, size=11, bold=True, color="1E293B")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    sub_header_font = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10, bold=False, color="0F172A")
    bold_data_font = Font(name=font_family, size=10, bold=True, color="0F172A")
    auto_font = Font(name=font_family, size=10, bold=True, color="1E3A8A")
    note_font = Font(name=font_family, size=9, italic=True, color="475569")

    # Fills
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_id_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")       # Dark Blue
    header_pillar_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")   # Teal
    header_dash_fill = PatternFill(start_color="6D28D9", end_color="6D28D9", fill_type="solid")     # Purple
    header_qg_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")       # Amber/Orange
    header_meta_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")     # Slate

    sub_input_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")       # Blue badge for Text Input
    sub_drop_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")        # Green badge for Dropdown
    sub_auto_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")        # Indigo badge for Formula

    zebra_even_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    zebra_odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    auto_col_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")        # Soft blue for auto-calculated cols
    section_banner_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    # Status Fills for Conditional Formatting
    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    pass_font = Font(name=font_family, size=10, bold=True, color="166534")

    warn_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    warn_font = Font(name=font_family, size=10, bold=True, color="854D0E")

    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fail_font = Font(name=font_family, size=10, bold=True, color="991B1B")

    none_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    none_font = Font(name=font_family, size=10, italic=True, color="64748B")

    # Borders
    thin_side = Side(border_style="thin", color="CBD5E1")
    med_navy = Side(border_style="medium", color="1E293B")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_border = Border(left=thin_side, right=thin_side, top=med_navy, bottom=med_navy)

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # =========================================================================
    # SHEET 3: 03_DANH_MUC_CHUAN (Catalogs & Reference Standards)
    # =========================================================================
    ws_catalog.views.sheetView[0].showGridLines = True

    ws_catalog.merge_cells("A1:H1")
    ws_catalog["A1"] = "DANH MỤC CHUẨN HÓA (DÙNG CHO DROPDOWN LIST) & TIÊU CHÍ THAM CHIẾU QUY ĐỊNH OBSERVABILITY"
    ws_catalog["A1"].font = title_font
    ws_catalog["A1"].fill = navy_fill
    ws_catalog["A1"].alignment = align_center
    ws_catalog.row_dimensions[1].height = 36

    catalog_headers = [
        ("A", "Danh mục Tech Stack", 34),
        ("B", "Danh mục Môi trường", 24),
        ("C", "Phương thức Tích hợp", 34),
        ("D", "Trạng thái 3 Trục (Log/Trace/Metric)", 30),
        ("E", "Trạng thái 3 Level Dashboard", 28),
        ("F", "Đường truyền Telemetry", 34),
        ("G", "Metric SDK Active (=1)", 26),
    ]

    ws_catalog.row_dimensions[3].height = 28
    for col_letter, text, width in catalog_headers:
        cell = ws_catalog[f"{col_letter}3"]
        cell.value = text
        cell.font = header_font
        cell.fill = header_id_fill
        cell.alignment = align_center
        cell.border = header_border
        ws_catalog.column_dimensions[col_letter].width = width

    catalog_data = {
        "A": [
            ".NET 8 (SDK hỗ trợ chính thức)",
            ".NET 6 / .NET 7",
            ".NET Framework (Legacy)",
            "Node.js / NestJS",
            "Python (FastAPI / Django)",
            "Go (Golang)",
            "Java / Spring Boot",
            "Khác"
        ],
        "B": [
            "Production",
            "Staging / UAT",
            "Development / SIT",
            "Chưa triển khai"
        ],
        "C": [
            "Đã cài SDK ISC.Observability (.NET 8)",
            "Cấu hình OTel Native (Biến môi trường)",
            "Đang tích hợp",
            "Chưa tích hợp"
        ],
        "D": [
            "Đã hoàn thành (Đạt chuẩn)",
            "Đã gắn (Chưa đạt chuẩn)",
            "Đang triển khai",
            "Chưa triển khai"
        ],
        "E": [
            "Đã có (Đạt chuẩn)",
            "Đang xây dựng",
            "Chưa có"
        ],
        "F": [
            "Qua OTel Collector (Đúng chuẩn)",
            "Đẩy thẳng Kafka / ES (Vi phạm BLOCKER)",
            "Chỉ mới xuất Console",
            "Chưa cấu hình"
        ],
        "G": [
            "Đã ghi nhận (= 1)",
            "Chưa ghi nhận (= 0)",
            "N/A (Ngoài .NET 8)"
        ]
    }

    for col_letter, items in catalog_data.items():
        for idx, val in enumerate(items, start=4):
            c = ws_catalog[f"{col_letter}{idx}"]
            c.value = val
            c.font = data_font
            c.alignment = align_left
            c.border = cell_border
            ws_catalog.row_dimensions[idx].height = 22

    # Reference table on Sheet 3 (Rows 14+)
    ws_catalog.merge_cells("A14:G14")
    ws_catalog["A14"] = "BẢNG TÓM TẮT TIÊU CHÍ ĐÁNH GIÁ TUÂN THỦ (TRÍCH XUẤT TỪ QUY ĐỊNH ISC OBSERVABILITY v1.0)"
    ws_catalog["A14"].font = header_font
    ws_catalog["A14"].fill = header_pillar_fill
    ws_catalog["A14"].alignment = align_left
    ws_catalog.row_dimensions[14].height = 28

    ref_rows = [
        ("Hạng mục kiểm tra", "Tiêu chí công nhận 'Đạt chuẩn'", "Cơ chế hỗ trợ trong SDK (.NET 8)"),
        ("Trục 1: Logging", "Xuất JSON có cấu trúc; đủ trường bắt buộc (timestamp, severity_text, service_name, environment, trace_id, span_id, correlation_id, message); biến nghiệp vụ dạng snake_case; che PII nhạy cảm.", "SDK tự động hóa 90% (Dev chỉ cần viết message và đặt tên biến nghiệp vụ snake_case)."),
        ("Trục 2: Tracing", "Lan truyền ngữ cảnh W3C Trace Context (traceparent) & B3; tự động tạo span cho HTTP/DB/Redis; tạo custom span cho tác vụ nội bộ > 200ms.", "SDK tự động gắn ASP.NET Core, HttpClient, EF Core, SqlClient, Redis, Quartz instrumentation."),
        ("Trục 3: Metrics", "Đầy đủ Golden Signals (RPS, Latency p50/p95/p99, Error Rate); phát metric tuân thủ observability.sdk.active = 1; nhãn (label) dạng snake_case và Low Cardinality.", "SDK tự động phát Golden Signals, Runtime metrics và observability.sdk.active = 1."),
        ("Dashboard Level 1", "Executive Overview (Dành cho Ban điều hành): Hiển thị Golden Signals gồm Traffic (RPS), Latency (p50/p95/p99), Error Rate (%), Saturation.", "Dữ liệu tự động đổ về từ OTel Collector sang Elasticsearch/Kibana hoặc SigNoz."),
        ("Dashboard Level 2", "Developer Deep-Dive (Dành cho Dev/Tech Lead): APM chi tiết, tra cứu Waterfall theo Trace ID, danh sách Slow Endpoints & Error Logs.", "Truy vết xuyên suốt qua trường trace_id đồng nhất giữa Log và Span."),
        ("Dashboard Level 3", "Infrastructure & Runtime (Dành cho DevOps/SRE): Giám sát CPU, Memory, GC Pause Time, ThreadPool Queue.", "SDK tự động thu thập qua OpenTelemetry Runtime Instrumentation."),
        ("Đường truyền OTLP", "Bắt buộc 100% đi qua OpenTelemetry Collector (Single OTLP Path - cổng 4318 HTTP Protobuf). Nghiêm cấm app đẩy trực tiếp vào Kafka/Elasticsearch.", "SDK cấu hình mặc định OTLP HTTP Protobuf qua biến môi trường OTEL_EXPORTER_OTLP_ENDPOINT.")
    ]

    for idx, (col1, col2, col3) in enumerate(ref_rows, start=15):
        ws_catalog.merge_cells(f"B{idx}:D{idx}")
        ws_catalog.merge_cells(f"E{idx}:G{idx}")
        ws_catalog[f"A{idx}"] = col1
        ws_catalog[f"B{idx}"] = col2
        ws_catalog[f"E{idx}"] = col3

        is_header = (idx == 15)
        for col_c in ["A", "B", "C", "D", "E", "F", "G"]:
            cell = ws_catalog[f"{col_c}{idx}"]
            cell.border = header_border if is_header else cell_border
            cell.font = bold_data_font if is_header else data_font
            cell.fill = section_banner_fill if is_header else zebra_odd_fill
            cell.alignment = align_left
        ws_catalog.row_dimensions[idx].height = 34 if not is_header else 26

    # =========================================================================
    # SHEET 2: 02_TRACKING_SERVICES (Main Interactive Tracking Form)
    # =========================================================================
    ws_tracking.views.sheetView[0].showGridLines = True
    MAX_ROWS = 200
    START_ROW = 6
    END_ROW = START_ROW + MAX_ROWS - 1  # Row 6 to 205

    # Title Banner
    ws_tracking.merge_cells("A1:V1")
    ws_tracking["A1"] = "BIỂU MẪU THEO DÕI MỨC ĐỘ TUÂN THỦ TIÊU CHUẨN OBSERVABILITY (3 TRỤC & 3 LEVEL DASHBOARD)"
    ws_tracking["A1"].font = title_font
    ws_tracking["A1"].fill = navy_fill
    ws_tracking["A1"].alignment = align_center
    ws_tracking.row_dimensions[1].height = 34

    ws_tracking.merge_cells("A2:V2")
    ws_tracking["A2"] = (
        "Hướng dẫn: Các cột nhãn [NHẬP TEXT] do đơn vị tự điền | Các cột nhãn [CHỌN DANH MỤC] bấm mũi tên chọn từ danh sách chuẩn | "
        "Các cột nhãn [TỰ ĐỘNG TÍNH] có nền xanh nhạt do hệ thống tự động đánh giá."
    )
    ws_tracking["A2"].font = subtitle_font
    ws_tracking["A2"].fill = navy_fill
    ws_tracking["A2"].alignment = align_center
    ws_tracking.row_dimensions[2].height = 22

    # Group Headers (Row 3)
    groups = [
        ("A3:G3", "1. THÔNG TIN ĐỊNH DANH SERVICE & TECH STACK", header_id_fill),
        ("H3:K3", "2. TRẠNG THÁI 3 TRỤC OBSERVABILITY (3 PILLARS)", header_pillar_fill),
        ("L3:O3", "3. TRẠNG THÁI 3 LEVEL DASHBOARD (MỤC 8 QUY ĐỊNH)", header_dash_fill),
        ("P3:R3", "4. HẠ TẦNG OTLP & NGHIỆM THU QUALITY GATE 2", header_qg_fill),
        ("S3:V3", "5. THÔNG TIN QUẢN TRỊ & LIÊN KẾT", header_meta_fill),
    ]
    ws_tracking.row_dimensions[3].height = 26
    for cell_range, title, fill in groups:
        ws_tracking.merge_cells(cell_range)
        start_col_letter = cell_range.split(":")[0][0]
        top_left_cell = ws_tracking[cell_range.split(":")[0]]
        top_left_cell.value = title
        top_left_cell.font = header_font
        top_left_cell.alignment = align_center
        # Apply border/fill to all merged cells
        col_start, row_s, col_end, row_e = openpyxl.utils.range_boundaries(cell_range)
        for c_idx in range(col_start, col_end + 1):
            c = ws_tracking.cell(row=row_s, column=c_idx)
            c.fill = fill
            c.border = header_border

    # Column Definitions (Row 4 = Header Name, Row 5 = Input Type Badge)
    columns_spec = [
        # Col, Title, Badge, BadgeFill, Width, GroupFill
        ("A", "STT", "[TỰ ĐỘNG]", sub_auto_fill, 7, header_id_fill),
        ("B", "Tên Hệ thống / Dự án", "[NHẬP TEXT]", sub_input_fill, 24, header_id_fill),
        ("C", "Tên Service\n(Chuẩn kebab-case-svc)", "[NHẬP TEXT]", sub_input_fill, 26, header_id_fill),
        ("D", "Môi trường\n(Environment)", "[CHỌN DANH MỤC]", sub_drop_fill, 18, header_id_fill),
        ("E", "Tech Stack của Service", "[CHỌN DANH MỤC]", sub_drop_fill, 28, header_id_fill),
        ("F", "Khả năng hỗ trợ của SDK\n(Hiện chỉ hỗ trợ .NET 8)", "[TỰ ĐỘNG TÍNH]", sub_auto_fill, 30, header_id_fill),
        ("G", "Phương thức Tích hợp", "[CHỌN DANH MỤC]", sub_drop_fill, 28, header_id_fill),

        ("H", "Trục 1: LOGGING\n(JSON, snake_case, Mask PII)", "[CHỌN DANH MỤC]", sub_drop_fill, 25, header_pillar_fill),
        ("I", "Trục 2: TRACING\n(W3C/B3, Auto & Custom Span)", "[CHỌN DANH MỤC]", sub_drop_fill, 25, header_pillar_fill),
        ("J", "Trục 3: METRICS\n(Golden Signals, Low Card.)", "[CHỌN DANH MỤC]", sub_drop_fill, 25, header_pillar_fill),
        ("K", "Tổng hợp 3 Trục\n(Full 3 Pillars)", "[TỰ ĐỘNG TÍNH]", sub_auto_fill, 24, header_pillar_fill),

        ("L", "Dashboard Level 1\nExecutive Overview", "[CHỌN DANH MỤC]", sub_drop_fill, 22, header_dash_fill),
        ("M", "Dashboard Level 2\nDeveloper Deep-Dive", "[CHỌN DANH MỤC]", sub_drop_fill, 22, header_dash_fill),
        ("N", "Dashboard Level 3\nInfra & Runtime", "[CHỌN DANH MỤC]", sub_drop_fill, 22, header_dash_fill),
        ("O", "Tổng hợp Dashboard\n(Full 3 Levels)", "[TỰ ĐỘNG TÍNH]", sub_auto_fill, 24, header_dash_fill),

        ("P", "Đường truyền Telemetry\n(Single OTLP Path)", "[CHỌN DANH MỤC]", sub_drop_fill, 28, header_qg_fill),
        ("Q", "Tín hiệu SDK Active\n(observability.sdk.active=1)", "[CHỌN DANH MỤC]", sub_drop_fill, 24, header_qg_fill),
        ("R", "KẾT QUẢ TUÂN THỦ\n(QUALITY GATE 2)", "[TỰ ĐỘNG TÍNH]", sub_auto_fill, 26, header_qg_fill),

        ("S", "Link Dashboard\n(Kibana / SigNoz / Grafana)", "[NHẬP TEXT]", sub_input_fill, 30, header_meta_fill),
        ("T", "Tech Lead / PIC", "[NHẬP TEXT]", sub_input_fill, 20, header_meta_fill),
        ("U", "Ngày cập nhật /\nDeadline hoàn thành", "[NHẬP TEXT]", sub_input_fill, 18, header_meta_fill),
        ("V", "Ghi chú / Vướng mắc", "[NHẬP TEXT]", sub_input_fill, 32, header_meta_fill),
    ]

    ws_tracking.row_dimensions[4].height = 32
    ws_tracking.row_dimensions[5].height = 20

    for col_letter, title, badge, badge_fill, width, group_fill in columns_spec:
        ws_tracking.column_dimensions[col_letter].width = width

        c_header = ws_tracking[f"{col_letter}4"]
        c_header.value = title
        c_header.font = header_font
        c_header.fill = group_fill
        c_header.alignment = align_center
        c_header.border = header_border

        c_badge = ws_tracking[f"{col_letter}5"]
        c_badge.value = badge
        c_badge.font = sub_header_font
        c_badge.fill = badge_fill
        c_badge.alignment = align_center
        c_badge.border = cell_border

    # Freeze Panes & AutoFilter
    ws_tracking.freeze_panes = "D6"
    ws_tracking.auto_filter.ref = f"A5:V{END_ROW}"

    # Sample pre-filled rows to demonstrate functionality immediately
    sample_rows = [
        {
            "B": "Hệ thống Thanh toán (Core Payment)",
            "C": "payment-svc",
            "D": "Production",
            "E": ".NET 8 (SDK hỗ trợ chính thức)",
            "G": "Đã cài SDK ISC.Observability (.NET 8)",
            "H": "Đã hoàn thành (Đạt chuẩn)",
            "I": "Đã hoàn thành (Đạt chuẩn)",
            "J": "Đã hoàn thành (Đạt chuẩn)",
            "L": "Đã có (Đạt chuẩn)",
            "M": "Đã có (Đạt chuẩn)",
            "N": "Đã có (Đạt chuẩn)",
            "P": "Qua OTel Collector (Đúng chuẩn)",
            "Q": "Đã ghi nhận (= 1)",
            "S": "https://kibana.isc.internal/app/dashboards#/view/payment-svc",
            "T": "Nguyễn Văn A",
            "U": "2026-09-29",
            "V": "Mẫu: Đã đạt chuẩn toàn diện Quality Gate 2 (Có thể xóa/ghi đè dòng mẫu này)"
        },
        {
            "B": "Hệ thống Đơn hàng (E-Commerce)",
            "C": "order-svc",
            "D": "Staging / UAT",
            "E": ".NET 8 (SDK hỗ trợ chính thức)",
            "G": "Đã cài SDK ISC.Observability (.NET 8)",
            "H": "Đã hoàn thành (Đạt chuẩn)",
            "I": "Đã hoàn thành (Đạt chuẩn)",
            "J": "Đang triển khai",
            "L": "Đã có (Đạt chuẩn)",
            "M": "Đã có (Đạt chuẩn)",
            "N": "Đang xây dựng",
            "P": "Qua OTel Collector (Đúng chuẩn)",
            "Q": "Đã ghi nhận (= 1)",
            "S": "https://kibana.isc.internal/app/dashboards#/view/order-svc",
            "T": "Trần Thị B",
            "U": "2026-10-05",
            "V": "Mẫu: Đang hoàn thiện custom metrics & Dashboard Layer 3"
        },
        {
            "B": "Hệ thống Thông báo (Notification Hub)",
            "C": "notification-svc",
            "D": "Production",
            "E": "Node.js / NestJS",
            "G": "Cấu hình OTel Native (Biến môi trường)",
            "H": "Đã hoàn thành (Đạt chuẩn)",
            "I": "Đã hoàn thành (Đạt chuẩn)",
            "J": "Đã hoàn thành (Đạt chuẩn)",
            "L": "Đã có (Đạt chuẩn)",
            "M": "Đã có (Đạt chuẩn)",
            "N": "Đã có (Đạt chuẩn)",
            "P": "Qua OTel Collector (Đúng chuẩn)",
            "Q": "N/A (Ngoài .NET 8)",
            "S": "https://kibana.isc.internal/app/dashboards#/view/notification-svc",
            "T": "Lê Văn C",
            "U": "2026-09-28",
            "V": "Mẫu: Service Node.js cấu hình qua biến môi trường chuẩn OTLP"
        },
        {
            "B": "Hệ thống Đối soát Cũ (Legacy Billing)",
            "C": "billing-worker-svc",
            "D": "Development / SIT",
            "E": ".NET 6 / .NET 7",
            "G": "Đang tích hợp",
            "H": "Đã gắn (Chưa đạt chuẩn)",
            "I": "Chưa triển khai",
            "J": "Chưa triển khai",
            "L": "Chưa có",
            "M": "Chưa có",
            "N": "Chưa có",
            "P": "Đẩy thẳng Kafka / ES (Vi phạm BLOCKER)",
            "Q": "Chưa ghi nhận (= 0)",
            "S": "",
            "T": "Phạm Văn D",
            "U": "2026-10-20",
            "V": "Mẫu: Đang vi phạm đẩy log thẳng vào Kafka, chờ nâng cấp lên .NET 8"
        },
        {
            "B": "Hệ thống Gợi ý AI (Recommendation)",
            "C": "ai-recommend-svc",
            "D": "Chưa triển khai",
            "E": "Python (FastAPI / Django)",
            "G": "Chưa tích hợp",
            "H": "Chưa triển khai",
            "I": "Chưa triển khai",
            "J": "Chưa triển khai",
            "L": "Chưa có",
            "M": "Chưa có",
            "N": "Chưa có",
            "P": "Chưa cấu hình",
            "Q": "N/A (Ngoài .NET 8)",
            "S": "",
            "T": "Hoàng Văn E",
            "U": "2026-11-01",
            "V": "Mẫu: Dự án mới khởi tạo, chưa tích hợp"
        }
    ]

    # Populate rows 6 to 205 with formulas and formatting
    for r_idx in range(START_ROW, END_ROW + 1):
        ws_tracking.row_dimensions[r_idx].height = 24
        row_fill = zebra_even_fill if r_idx % 2 == 0 else zebra_odd_fill

        # Formula A: Auto Row Number when Service Name (C) is entered
        ws_tracking[f"A{r_idx}"] = f'=IF(COUNTA(B{r_idx}:C{r_idx})=0, "", COUNTA($C$6:C{r_idx}))'

        # Formula F: SDK Compatibility based on Tech Stack (E)
        ws_tracking[f"F{r_idx}"] = (
            f'=IF(E{r_idx}="", "", '
            f'IF(E{r_idx}=".NET 8 (SDK hỗ trợ chính thức)", "Hỗ trợ chuẩn SDK (.NET 8)", '
            f'"Ngoài phạm vi SDK (Dùng OTel Native)"))'
        )

        # Formula K: 3 Pillars Evaluation (H, I, J)
        ws_tracking[f"K{r_idx}"] = (
            f'=IF(COUNTA(H{r_idx}:J{r_idx})=0, "", '
            f'IF(COUNTIF(H{r_idx}:J{r_idx}, "Đã hoàn thành (Đạt chuẩn)")=3, "ĐỦ 3 TRỤC (100%)", '
            f'IF(COUNTIF(H{r_idx}:J{r_idx}, "Chưa triển khai")=3, "CHƯA GẮN (0%)", '
            f'"THIẾU / CHƯA ĐẠT (" & COUNTIF(H{r_idx}:J{r_idx}, "Đã hoàn thành (Đạt chuẩn)") & "/3)")))'
        )

        # Formula O: 3 Dashboard Levels Evaluation (L, M, N)
        ws_tracking[f"O{r_idx}"] = (
            f'=IF(COUNTA(L{r_idx}:N{r_idx})=0, "", '
            f'IF(COUNTIF(L{r_idx}:N{r_idx}, "Đã có (Đạt chuẩn)")=3, "ĐỦ 3 LEVEL DASHBOARD", '
            f'IF(COUNTIF(L{r_idx}:N{r_idx}, "Chưa có")=3, "CHƯA CÓ DASHBOARD", '
            f'"THIẾU DASHBOARD (" & COUNTIF(L{r_idx}:N{r_idx}, "Đã có (Đạt chuẩn)") & "/3)")))'
        )

        # Formula R: Overall Quality Gate 2 Compliance
        ws_tracking[f"R{r_idx}"] = (
            f'=IF(C{r_idx}="", "", '
            f'IF(P{r_idx}="Đẩy thẳng Kafka / ES (Vi phạm BLOCKER)", "VI PHẠM (BLOCKER)", '
            f'IF(AND(K{r_idx}="ĐỦ 3 TRỤC (100%)", O{r_idx}="ĐỦ 3 LEVEL DASHBOARD", P{r_idx}="Qua OTel Collector (Đúng chuẩn)"), "ĐẠT CHUẨN QG2 (PASS)", '
            f'IF(AND(K{r_idx}="CHƯA GẮN (0%)", O{r_idx}="CHƯA CÓ DASHBOARD"), "CHƯA TRIỂN KHAI", '
            f'"ĐANG HOÀN THIỆN (PENDING)"))))'
        )

        # Fill sample data if within sample rows
        sample_idx = r_idx - START_ROW
        if sample_idx < len(sample_rows):
            for col_k, val in sample_rows[sample_idx].items():
                ws_tracking[f"{col_k}{r_idx}"] = val

        # Apply borders, fonts, and alignments
        center_cols = {"A", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "U"}
        auto_cols = {"A", "F", "K", "O", "R"}

        for col_idx in range(1, 23):
            col_l = get_column_letter(col_idx)
            cell = ws_tracking[f"{col_l}{r_idx}"]
            cell.border = cell_border
            cell.alignment = align_center if col_l in center_cols else align_left
            if col_l in auto_cols:
                cell.fill = auto_col_fill
                cell.font = auto_font
            else:
                cell.fill = row_fill
                cell.font = bold_data_font if col_l == "C" else data_font

    # =========================================================================
    # DATA VALIDATIONS (Dropdown Menus on Sheet 2 referencing Sheet 3)
    # =========================================================================
    validations = [
        ("D", "'03_DANH_MUC_CHUAN'!$B$4:$B$7", "Chọn Môi trường", "Vui lòng chọn môi trường từ danh sách"),
        ("E", "'03_DANH_MUC_CHUAN'!$A$4:$A$11", "Chọn Tech Stack", "Hiện tại SDK chỉ hỗ trợ chính thức .NET 8"),
        ("G", "'03_DANH_MUC_CHUAN'!$C$4:$C$7", "Chọn Phương thức tích hợp", "Chọn phương thức tích hợp"),
        ("H", "'03_DANH_MUC_CHUAN'!$D$4:$D$7", "Trạng thái Logging", "Chọn trạng thái triển khai trục Logging"),
        ("I", "'03_DANH_MUC_CHUAN'!$D$4:$D$7", "Trạng thái Tracing", "Chọn trạng thái triển khai trục Tracing"),
        ("J", "'03_DANH_MUC_CHUAN'!$D$4:$D$7", "Trạng thái Metrics", "Chọn trạng thái triển khai trục Metrics"),
        ("L", "'03_DANH_MUC_CHUAN'!$E$4:$E$6", "Dashboard Level 1", "Chọn trạng thái Dashboard Level 1 (Executive)"),
        ("M", "'03_DANH_MUC_CHUAN'!$E$4:$E$6", "Dashboard Level 2", "Chọn trạng thái Dashboard Level 2 (Developer)"),
        ("N", "'03_DANH_MUC_CHUAN'!$E$4:$E$6", "Dashboard Level 3", "Chọn trạng thái Dashboard Level 3 (Infra & Runtime)"),
        ("P", "'03_DANH_MUC_CHUAN'!$F$4:$F$7", "Đường truyền Telemetry", "Chọn kiến trúc đường truyền OTLP"),
        ("Q", "'03_DANH_MUC_CHUAN'!$G$4:$G$6", "Tín hiệu SDK Active", "Kiểm tra metric observability.sdk.active = 1"),
    ]

    for col_l, formula_ref, prompt_title, prompt_body in validations:
        dv = DataValidation(
            type="list",
            formula1=f"={formula_ref}",
            allow_blank=True,
            showDropDown=False,  # In openpyxl, False means SHOW the dropdown arrow in Excel!
            promptTitle=prompt_title,
            prompt=prompt_body
        )
        ws_tracking.add_data_validation(dv)
        dv.add(f"{col_l}{START_ROW}:{col_l}{END_ROW}")

    # =========================================================================
    # CONDITIONAL FORMATTING ON SHEET 2
    # =========================================================================
    full_range = f"F{START_ROW}:R{END_ROW}"

    green_values = [
        '"Đã hoàn thành (Đạt chuẩn)"',
        '"Đã có (Đạt chuẩn)"',
        '"ĐỦ 3 TRỤC (100%)"',
        '"ĐỦ 3 LEVEL DASHBOARD"',
        '"Qua OTel Collector (Đúng chuẩn)"',
        '"Đã ghi nhận (= 1)"',
        '"ĐẠT CHUẨN QG2 (PASS)"',
        '"Hỗ trợ chuẩn SDK (.NET 8)"',
        '"Đã cài SDK ISC.Observability (.NET 8)"'
    ]
    for val in green_values:
        ws_tracking.conditional_formatting.add(
            full_range,
            CellIsRule(operator="equal", formula=[val], stopIfTrue=True, fill=pass_fill, font=pass_font)
        )

    red_values = [
        '"Đẩy thẳng Kafka / ES (Vi phạm BLOCKER)"',
        '"VI PHẠM (BLOCKER)"',
        '"Chưa ghi nhận (= 0)"'
    ]
    for val in red_values:
        ws_tracking.conditional_formatting.add(
            full_range,
            CellIsRule(operator="equal", formula=[val], stopIfTrue=True, fill=fail_fill, font=fail_font)
        )

    yellow_values = [
        '"Đã gắn (Chưa đạt chuẩn)"',
        '"Đang triển khai"',
        '"Đang xây dựng"',
        '"ĐANG HOÀN THIỆN (PENDING)"',
        '"Ngoài phạm vi SDK (Dùng OTel Native)"',
        '"Chỉ mới xuất Console"'
    ]
    for val in yellow_values:
        ws_tracking.conditional_formatting.add(
            full_range,
            CellIsRule(operator="equal", formula=[val], stopIfTrue=True, fill=warn_fill, font=warn_font)
        )

    gray_values = [
        '"Chưa triển khai"',
        '"Chưa có"',
        '"CHƯA GẮN (0%)"',
        '"CHƯA CÓ DASHBOARD"',
        '"CHƯA TRIỂN KHAI"',
        '"Chưa cấu hình"',
        '"Chưa tích hợp"'
    ]
    for val in gray_values:
        ws_tracking.conditional_formatting.add(
            full_range,
            CellIsRule(operator="equal", formula=[val], stopIfTrue=True, fill=none_fill, font=none_font)
        )

    # =========================================================================
    # SHEET 1: 01_TONG_QUAN (Executive Summary & Automatic KPI Dashboard)
    # =========================================================================
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.merge_cells("A1:F1")
    ws_summary["A1"] = "BẢNG TỔNG HỢP TIẾN ĐỘ TUÂN THỦ OBSERVABILITY TOÀN HỆ THỐNG"
    ws_summary["A1"].font = title_font
    ws_summary["A1"].fill = navy_fill
    ws_summary["A1"].alignment = align_center
    ws_summary.row_dimensions[1].height = 36

    ws_summary.merge_cells("A2:F2")
    ws_summary["A2"] = "Số liệu tự động tổng hợp theo thời gian thực từ Sheet '02_TRACKING_SERVICES' (Dựa trên Quy định Observability v1.0 & SDK v1.4.3)"
    ws_summary["A2"].font = subtitle_font
    ws_summary["A2"].fill = navy_fill
    ws_summary["A2"].alignment = align_center
    ws_summary.row_dimensions[2].height = 22

    summary_cols = [("A", 6), ("B", 44), ("C", 20), ("D", 20), ("E", 38), ("F", 6)]
    for col_l, w in summary_cols:
        ws_summary.column_dimensions[col_l].width = w

    def add_summary_section_header(row_idx, title, fill_color):
        ws_summary.merge_cells(f"B{row_idx}:E{row_idx}")
        ws_summary[f"B{row_idx}"] = title
        ws_summary.row_dimensions[row_idx].height = 28
        for c_l in ["B", "C", "D", "E"]:
            cell = ws_summary[f"{c_l}{row_idx}"]
            cell.font = header_font
            cell.fill = fill_color
            cell.border = header_border
            if c_l == "B":
                cell.alignment = align_left

    def add_summary_table_header(row_idx, col_b, col_c, col_d, col_e):
        ws_summary.row_dimensions[row_idx].height = 24
        headers = [("B", col_b), ("C", col_c), ("D", col_d), ("E", col_e)]
        for c_l, h_text in headers:
            cell = ws_summary[f"{c_l}{row_idx}"]
            cell.value = h_text
            cell.font = bold_data_font
            cell.fill = section_banner_fill
            cell.border = header_border
            cell.alignment = align_center if c_l in ["C", "D"] else align_left

    def add_summary_row(row_idx, label, count_formula, pct_formula, note, highlight_fill=None):
        ws_summary.row_dimensions[row_idx].height = 24
        ws_summary[f"B{row_idx}"] = label
        ws_summary[f"C{row_idx}"] = count_formula
        ws_summary[f"D{row_idx}"] = pct_formula
        ws_summary[f"E{row_idx}"] = note

        for c_l in ["B", "C", "D", "E"]:
            cell = ws_summary[f"{c_l}{row_idx}"]
            cell.border = cell_border
            cell.fill = highlight_fill if highlight_fill else zebra_odd_fill
            cell.font = bold_data_font if highlight_fill else data_font
            if c_l in ["C", "D"]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
        ws_summary[f"D{row_idx}"].number_format = "0.0%"

    # --- BLOCK 1: TỔNG QUAN CHUNG & NGHIỆM THU QUALITY GATE 2 ---
    add_summary_section_header(4, "I. TỔNG QUAN TUÂN THỦ & NGHIỆM THU QUALITY GATE 2", header_id_fill)
    add_summary_table_header(5, "Chỉ tiêu thống kê", "Số lượng Service", "Tỷ lệ (%)", "Ghi chú / Ý nghĩa")

    total_svc_ref = "$C$6"
    add_summary_row(
        6,
        "Tổng số Services đang theo dõi",
        f"=COUNTA('02_TRACKING_SERVICES'!C{START_ROW}:C{END_ROW})",
        "=IF(C6>0, 1, 0)",
        "Tổng số microservice có khai báo tên trong bảng Tracking",
        highlight_fill=auto_col_fill
    )
    add_summary_row(
        7,
        "1. ĐẠT CHUẨN QUALITY GATE 2 (PASS - Sẵn sàng Go-Live)",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!R{START_ROW}:R{END_ROW}, "ĐẠT CHUẨN QG2 (PASS)")',
        f"=IF({total_svc_ref}>0, C7/{total_svc_ref}, 0)",
        "Đủ 3 trục + Đủ 3 tầng Dashboard + Qua OTel Collector chuẩn",
        highlight_fill=pass_fill
    )
    add_summary_row(
        8,
        "2. ĐANG HOÀN THIỆN (PENDING - Thiếu trục hoặc Dashboard)",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!R{START_ROW}:R{END_ROW}, "ĐANG HOÀN THIỆN (PENDING)")',
        f"=IF({total_svc_ref}>0, C8/{total_svc_ref}, 0)",
        "Đã triển khai một phần nhưng chưa đủ điều kiện qua QG2",
        highlight_fill=warn_fill
    )
    add_summary_row(
        9,
        "3. VI PHẠM KIẾN TRÚC (BLOCKER - Đẩy thẳng Kafka/ES)",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!R{START_ROW}:R{END_ROW}, "VI PHẠM (BLOCKER)")',
        f"=IF({total_svc_ref}>0, C9/{total_svc_ref}, 0)",
        "Bắt buộc chặn release (Vi phạm mục 7.3 Single OTLP Path)",
        highlight_fill=fail_fill
    )
    add_summary_row(
        10,
        "4. CHƯA TRIỂN KHAI (NOT STARTED)",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!R{START_ROW}:R{END_ROW}, "CHƯA TRIỂN KHAI")',
        f"=IF({total_svc_ref}>0, C10/{total_svc_ref}, 0)",
        "Service chưa gắn bất kỳ trục hay dashboard nào",
        highlight_fill=none_fill
    )

    # --- BLOCK 2: THỐNG KÊ THEO 3 TRỤC OBSERVABILITY ---
    add_summary_section_header(12, "II. TIẾN ĐỘ TÍCH HỢP 3 TRỤC OBSERVABILITY (LOGGING - TRACING - METRICS)", header_pillar_fill)
    add_summary_table_header(13, "Trụ cột Observability", "Số Service Đạt chuẩn", "Tỷ lệ Hoàn thành", "Yêu cầu cốt lõi trong Quy định")

    add_summary_row(
        14,
        "Trục 1: LOGGING (Structured JSON, snake_case, PII Masking)",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!H{START_ROW}:H{END_ROW}, "Đã hoàn thành (Đạt chuẩn)")',
        f"=IF({total_svc_ref}>0, C14/{total_svc_ref}, 0)",
        "Mục 4: JSON chuẩn, không text tự do, không lộ PII"
    )
    add_summary_row(
        15,
        "Trục 2: TRACING (W3C/B3 Context, Auto & Custom Spans)",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!I{START_ROW}:I{END_ROW}, "Đã hoàn thành (Đạt chuẩn)")',
        f"=IF({total_svc_ref}>0, C15/{total_svc_ref}, 0)",
        "Mục 5: Truyền traceparent liên tục, khớp Trace ID với Log"
    )
    add_summary_row(
        16,
        "Trục 3: METRICS (Golden Signals, Low Cardinality)",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!J{START_ROW}:J{END_ROW}, "Đã hoàn thành (Đạt chuẩn)")',
        f"=IF({total_svc_ref}>0, C16/{total_svc_ref}, 0)",
        "Mục 6: RPS, Latency p50/p95/p99, Error Rate, SDK Active"
    )
    add_summary_row(
        17,
        "TỔNG SỐ SERVICE ĐÃ GẮN ĐẦY ĐỦ CẢ 3 TRỤC (FULL 3 PILLARS)",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!K{START_ROW}:K{END_ROW}, "ĐỦ 3 TRỤC (100%)")',
        f"=IF({total_svc_ref}>0, C17/{total_svc_ref}, 0)",
        "Hoàn thành đồng thời cả 3 trục Logging + Tracing + Metrics",
        highlight_fill=pass_fill
    )

    # --- BLOCK 3: THỐNG KÊ THEO 3 LEVEL DASHBOARD ---
    add_summary_section_header(19, "III. TIẾN ĐỘ XÂY DỰNG 3 LEVEL DASHBOARD (MỤC 8 QUY ĐỊNH)", header_dash_fill)
    add_summary_table_header(20, "Cấp độ Dashboard", "Số Service Đã có", "Tỷ lệ Hoàn thành", "Đối tượng sử dụng & Nội dung bắt buộc")

    add_summary_row(
        21,
        "Level 1: Executive Overview Dashboard",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!L{START_ROW}:L{END_ROW}, "Đã có (Đạt chuẩn)")',
        f"=IF({total_svc_ref}>0, C21/{total_svc_ref}, 0)",
        "Ban điều hành: Golden Signals (Traffic, Latency, Error, Saturation)"
    )
    add_summary_row(
        22,
        "Level 2: Developer Deep-Dive Dashboard",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!M{START_ROW}:M{END_ROW}, "Đã có (Đạt chuẩn)")',
        f"=IF({total_svc_ref}>0, C22/{total_svc_ref}, 0)",
        "Developer / Tech Lead: APM, truy vết theo Trace ID"
    )
    add_summary_row(
        23,
        "Level 3: Infrastructure & Runtime Dashboard",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!N{START_ROW}:N{END_ROW}, "Đã có (Đạt chuẩn)")',
        f"=IF({total_svc_ref}>0, C23/{total_svc_ref}, 0)",
        "DevOps / SRE: Memory, CPU, GC, ThreadPool"
    )
    add_summary_row(
        24,
        "TỔNG SỐ SERVICE ĐÃ ĐỦ CẢ 3 LEVEL DASHBOARD",
        f'=COUNTIF(\'02_TRACKING_SERVICES\'!O{START_ROW}:O{END_ROW}, "ĐỦ 3 LEVEL DASHBOARD")',
        f"=IF({total_svc_ref}>0, C24/{total_svc_ref}, 0)",
        "Đạt đủ cả 3 tầng Dashboard theo Mục 8",
        highlight_fill=pass_fill
    )

    # --- BLOCK 4: PHÂN BỔ THEO TECH STACK & KHẢ NĂNG HỖ TRỢ CỦA SDK ---
    add_summary_section_header(26, "IV. THỐNG KÊ THEO TECH STACK (ĐÁNH GIÁ ĐỘ PHỦ SDK .NET 8)", header_meta_fill)
    add_summary_table_header(27, "Tech Stack của Service", "Số lượng Service", "Tỷ trọng (%)", "Cơ chế Tích hợp Chuẩn")

    tech_stacks_summary = [
        (".NET 8 (SDK hỗ trợ chính thức)", "Cài trực tiếp NuGet package ISC.Observability v1.4.3", pass_fill),
        (".NET 6 / .NET 7", "Khuyến nghị nâng cấp .NET 8 hoặc dùng OTel Native", zebra_odd_fill),
        (".NET Framework (Legacy)", "Dùng OpenTelemetry .NET Framework SDK / OTel Native", zebra_odd_fill),
        ("Node.js / NestJS", "Dùng @opentelemetry/sdk-node + Biến môi trường chuẩn", zebra_odd_fill),
        ("Python (FastAPI / Django)", "Dùng opentelemetry-distro + Biến môi trường chuẩn", zebra_odd_fill),
        ("Go (Golang)", "Dùng go.opentelemetry.io/otel + Biến môi trường chuẩn", zebra_odd_fill),
        ("Java / Spring Boot", "Dùng opentelemetry-javaagent + Biến môi trường chuẩn", zebra_odd_fill),
        ("Khác", "Cấu hình xuất OTLP HTTP Protobuf về OTel Collector:4318", zebra_odd_fill),
    ]

    for idx, (stack_name, stack_note, fill_st) in enumerate(tech_stacks_summary, start=28):
        add_summary_row(
            idx,
            stack_name,
            f'=COUNTIF(\'02_TRACKING_SERVICES\'!E{START_ROW}:E{END_ROW}, "{stack_name}")',
            f"=IF({total_svc_ref}>0, C{idx}/{total_svc_ref}, 0)",
            stack_note,
            highlight_fill=fill_st if stack_name.startswith(".NET 8") else None
        )

    wb.save(filename)
    print(f"Successfully generated Excel file at: {filename}")

if __name__ == "__main__":
    create_observability_tracking_excel()
